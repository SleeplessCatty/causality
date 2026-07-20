from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "application-design-spec.md"
OUTPUT = ROOT / "docs" / "causal-network-application-design-spec.docx"

# standard_business_brief preset, with a named CJK font fallback override.
FONT_ASCII = "Hiragino Sans GB"
FONT_CJK = "Hiragino Sans GB"
BODY_COLOR = "172033"
MUTED_COLOR = "667085"
HEADING_BLUE = "2E74B5"
HEADING_DARK_BLUE = "1F4D78"
TABLE_HEADER_FILL = "F2F4F7"
TABLE_BORDER = "D7DCE3"
CODE_FILL = "F4F6F9"
ACCENT = "5B5BD6"
USABLE_DXA = 9360
TABLE_INDENT_DXA = 120


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_font(run, size=None, color=BODY_COLOR, bold=None, italic=None, name=FONT_ASCII):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, size, color=BODY_COLOR, bold=None):
    style.font.name = FONT_ASCII
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), FONT_ASCII)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_ASCII)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    style.font.size = Pt(size)
    style.font.color.rgb = rgb(color)
    if bold is not None:
        style.font.bold = bold


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), HEADING_BLUE)
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), FONT_ASCII)
    rfonts.set(qn("w:hAnsi"), FONT_ASCII)
    rfonts.set(qn("w:eastAsia"), FONT_CJK)
    rpr.append(rfonts)
    run.append(rpr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(
    r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\(https?://[^)]+\)|<https?://[^>]+>)"
)


def add_inline(paragraph, text: str):
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_font(run)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, size=9.5, color="384250", name="Menlo")
            run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        elif token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\((https?://[^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        else:
            url = token[1:-1]
            add_hyperlink(paragraph, url, url)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_font(run)


def shade_paragraph(paragraph, fill):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tcpr = cell._tc.get_or_add_tcPr()
    tc_mar = tcpr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tcpr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tblpr = table._tbl.tblPr
    borders = tblpr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), TABLE_BORDER)


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    tblpr = table._tbl.tblPr
    tblw = tblpr.first_child_found_in("w:tblW")
    tblw.set(qn("w:w"), str(sum(widths_dxa)))
    tblw.set(qn("w:type"), "dxa")
    tblind = tblpr.first_child_found_in("w:tblInd")
    if tblind is None:
        tblind = OxmlElement("w:tblInd")
        tblpr.append(tblind)
    tblind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tblind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.first_child_found_in("w:tcW")
            tcw.set(qn("w:w"), str(widths_dxa[idx]))
            tcw.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[idx] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
    set_table_borders(table)


def choose_widths(rows):
    columns = len(rows[0])
    if columns == 3:
        return [1900, 1450, 6010]
    if columns == 2:
        return [2400, 6960]
    lengths = []
    for idx in range(columns):
        lengths.append(max(4, max(len(row[idx]) for row in rows)))
    total = sum(lengths)
    raw = [max(1000, round(USABLE_DXA * length / total)) for length in lengths]
    raw[-1] += USABLE_DXA - sum(raw)
    return raw


def add_table(doc, rows):
    widths = choose_widths(rows)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    set_table_geometry(table, widths)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            if r_idx == 0:
                set_cell_shading(cell, TABLE_HEADER_FILL)
            add_inline(p, value)
            for run in p.runs:
                set_font(run, size=9.4, bold=(r_idx == 0))
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_font(run, size=9, color=MUTED_COLOR)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_font(run, size=9, color=MUTED_COLOR)


def add_num_definition(doc, kind, start_at=1):
    numbering = doc.part.numbering_part.element
    existing_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(existing_abs, default=-1) + 1
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    num_id = max(existing_num, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), str(start_at))
    lvl.append(start)
    numfmt = OxmlElement("w:numFmt")
    numfmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl.append(numfmt)
    lvltext = OxmlElement("w:lvlText")
    lvltext.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    lvl.append(lvltext)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    ppr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.append(spacing)
    lvl.append(ppr)
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), FONT_ASCII)
    rfonts.set(qn("w:hAnsi"), FONT_ASCII)
    rfonts.set(qn("w:eastAsia"), FONT_CJK)
    rpr.append(rfonts)
    lvl.append(rpr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    ppr = paragraph._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(num_id))
    numpr.append(ilvl)
    numpr.append(numid)
    ppr.append(numpr)


def add_list_item(doc, text, num_id):
    p = doc.add_paragraph(style="Normal")
    apply_numbering(p, num_id)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    p.paragraph_format.keep_together = True
    add_inline(p, text)


def add_code_block(doc, lines, language):
    if language == "mermaid":
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        shade_paragraph(p, CODE_FILL)
        relations = [
            "抽象事件  1 ─── N  具体事件",
            "抽象事件（原因） 1 ─── N  抽象因果关系  N ─── 1 抽象事件（结果）",
            "抽象因果关系  1 ─── N  具体因果依据",
            "具体事件（原因 / 结果） 1 ─── N  具体因果依据",
        ]
        for idx, line in enumerate(relations):
            if idx:
                p.add_run().add_break()
            run = p.add_run(line)
            set_font(run, size=9.5, color="384250", name="Menlo")
            run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        return
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.05
    shade_paragraph(p, CODE_FILL)
    for idx, line in enumerate(lines):
        if idx:
            p.add_run().add_break()
        run = p.add_run(line)
        set_font(run, size=8.8, color="384250", name="Menlo")
        run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    set_style_font(normal, 11, BODY_COLOR)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, HEADING_BLUE, 16, 8),
        ("Heading 2", 13, HEADING_BLUE, 12, 6),
        ("Heading 3", 12, HEADING_DARK_BLUE, 8, 4),
    ):
        style = doc.styles[name]
        set_style_font(style, size, color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = hp.add_run("产品设计规范")
    set_font(hr, size=9, color=MUTED_COLOR, bold=True)
    hp.paragraph_format.space_after = Pt(0)

    footer = section.footer
    fp = footer.paragraphs[0]
    add_page_number(fp)


def add_title_block(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run("因果网络应用设计规范")
    set_font(run, size=25, color=BODY_COLOR, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run("建立可追溯、可验证的事件与因果知识网络")
    set_font(run, size=13, color=MUTED_COLOR)

    for label, value in (
        ("版本", "V1.0"),
        ("状态", "可进入原型与开发"),
        ("日期", "2026-07-19"),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        label_run = p.add_run(f"{label}：")
        set_font(label_run, size=10.5, bold=True)
        value_run = p.add_run(value)
        set_font(value_run, size=10.5)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(8)


def render_markdown(doc, lines):
    index = 0
    bullet_num_id = add_num_definition(doc, "bullet")
    current_ordered_id = None
    in_ordered = False
    while index < len(lines):
        raw = lines[index]
        line = raw.rstrip()
        if not line:
            in_ordered = False
            current_ordered_id = None
            index += 1
            continue
        if line.startswith("```"):
            language = line[3:].strip()
            block = []
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                block.append(lines[index].rstrip("\n"))
                index += 1
            add_code_block(doc, block, language)
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|?\s*:?-+", lines[index + 1]):
            rows = []
            header = [x.strip() for x in line.strip("|").split("|")]
            rows.append(header)
            index += 2
            while index < len(lines) and lines[index].lstrip().startswith("|"):
                rows.append([x.strip() for x in lines[index].strip().strip("|").split("|")])
                index += 1
            add_table(doc, rows)
            continue
        heading = re.match(r"^(#{2,4})\s+(.+)$", line)
        if heading:
            level = min(len(heading.group(1)) - 1, 3)
            p = doc.add_paragraph(style=f"Heading {level}")
            add_inline(p, heading.group(2))
            index += 1
            continue
        bullet = re.match(r"^-\s+(.+)$", line)
        if bullet:
            add_list_item(doc, bullet.group(1), bullet_num_id)
            index += 1
            continue
        ordered = re.match(r"^\d+\.\s+(.+)$", line)
        if ordered:
            if not in_ordered:
                current_ordered_id = add_num_definition(doc, "decimal")
                in_ordered = True
            add_list_item(doc, ordered.group(1), current_ordered_id)
            index += 1
            continue
        if line.startswith("**") and line.endswith("**"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[2:-2])
            set_font(run, bold=True, color=HEADING_DARK_BLUE)
            index += 1
            continue
        paragraph_lines = [line]
        index += 1
        while index < len(lines):
            next_line = lines[index].rstrip()
            if not next_line or next_line.startswith(("#", "- ", "```", "|")) or re.match(r"^\d+\.\s+", next_line):
                break
            paragraph_lines.append(next_line)
            index += 1
        p = doc.add_paragraph(style="Normal")
        add_inline(p, " ".join(paragraph_lines))


def main():
    markdown = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)
    add_title_block(doc)
    # Skip the Markdown title and its three metadata lines.
    render_markdown(doc, markdown[6:])
    doc.core_properties.title = "因果网络应用设计规范"
    doc.core_properties.subject = "事件、因果关系、依据库与网络可视化设计"
    doc.core_properties.author = "Codex"
    doc.core_properties.keywords = "因果网络, 产品设计, 应用规范"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
